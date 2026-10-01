"""Build a synthetic ZIP of monthly production workbooks with deliberately
different layouts, used to test extract_production.py.

    python tests/make_sample_zip.py <out.zip>
"""
from __future__ import annotations

import calendar
import datetime as dt
import io
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font


def wb_bytes(wb) -> bytes:
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def val(seed: int, day: int, base: float) -> float:
    return round(base + ((seed * 37 + day * 13) % 50) - 25, 1)


def vertical_merged(year: int, month: int, extra_pickling: bool = False, unmerged: bool = False,
                    formula_total: bool = False) -> bytes:
    """Dates down column A; 2-row header CRM04/CRM06 x Input/Output/Target; Total/Avg rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = calendar.month_abbr[month] + str(year)[2:]
    ws["A1"] = f"CRM SAHIBABAD - DAILY PRODUCTION REPORT FOR {calendar.month_name[month].upper()} {year}"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:K1")
    mills = (["PICKLING"] if extra_pickling else []) + ["CRM04", "CRM06"]
    ws.cell(3, 1, "Date")
    ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
    col = 2
    for mill in mills:
        ws.cell(3, col, mill)
        if not unmerged:
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 2)
        for j, m in enumerate(["Input (MT)", "Output (MT)", "Target"]):
            ws.cell(4, col + j, m)
        col += 3
    ws.cell(3, col, "Total Output (MT)")
    ws.cell(3, col + 1, "Remarks")
    days = calendar.monthrange(year, month)[1]
    r = 5
    sums = [0.0] * (len(mills) * 3 + 1)
    for d in range(1, days + 1):
        ws.cell(r, 1, dt.datetime(year, month, d))
        ws.cell(r, 1).number_format = "dd-mm-yyyy"
        tot = 0.0
        for i, mill in enumerate(mills):
            inp, out, tgt = val(i, d, 1000), val(i + 5, d, 950), 1000
            ws.cell(r, 2 + 3 * i, inp)
            ws.cell(r, 3 + 3 * i, out)
            ws.cell(r, 4 + 3 * i, tgt)
            sums[3 * i] += inp
            sums[3 * i + 1] += out
            sums[3 * i + 2] += tgt
            tot += out
        ws.cell(r, 2 + 3 * len(mills), round(tot, 1))
        sums[-1] += round(tot, 1)
        r += 1
    # special cells
    ws.cell(7, 2, "-")                    # placeholder in a data cell (day 3 CRM/PICKLING input)
    sums[0] -= val(0, 3, 1000)
    ws.cell(8, 3, "1,234.5")              # number stored as text
    sums[1] += 1234.5 - val(5, 4, 950)
    ws.cell(9, 2 + 3 * len(mills) + 1, "Roll change delay 2 hrs")
    ws.row_dimensions[10].hidden = True
    ws.cell(r, 1, "TOTAL")
    for j, s in enumerate(sums):
        if formula_total:
            L = openpyxl.utils.get_column_letter(2 + j)
            ws.cell(r, 2 + j, f"=SUM({L}5:{L}{r - 1})")
        else:
            ws.cell(r, 2 + j, round(s, 1))
    if not formula_total:
        ws.cell(r, 2, round(sums[0] + 7, 1))   # deliberate mismatch -> reconciliation must flag it
    ws.cell(r + 1, 1, "AVERAGE")
    for j, s in enumerate(sums):
        ws.cell(r + 1, 2 + j, round(s / days, 2))
    ws.cell(r + 3, 1, "Note: Figures are provisional. Prepared by Production Planning.")
    ws.cell(r + 4, 1, "Monthly ABP target")
    ws.cell(r + 4, 3, 30000)
    # second sheet must be ignored
    ws2 = wb.create_sheet("Delays")
    ws2["A1"] = "Delay"
    ws2["B1"] = 999
    return wb_bytes(wb)


def horizontal(year: int, month: int) -> bytes:
    """Dates across row 4 (day numbers), Area merged down col A, measure in col B, Total col."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prod"
    ws["A1"] = f"Production Summary {calendar.month_abbr[month]}-{str(year)[2:]}"
    days = calendar.monthrange(year, month)[1]
    ws.cell(4, 1, "Mill")
    ws.cell(4, 2, "Parameter")
    for d in range(1, days + 1):
        ws.cell(4, 2 + d, d)
    ws.cell(4, 3 + days, "Total")
    r = 5
    for i, mill in enumerate(["CRM04", "CRM06", "BAF"]):
        ws.cell(r, 1, mill)
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=1)
        for j, meas in enumerate(["Input (MT)", "Output (MT)", "Target (MT)"]):
            ws.cell(r + j, 2, meas)
            s = 0
            for d in range(1, days + 1):
                v = val(i * 3 + j, d, 800)
                ws.cell(r + j, 2 + d, v)
                s += v
            ws.cell(r + j, 3 + days, round(s, 1))
        r += 3
    ws.cell(r, 1, "Grand Total Output")
    for d in range(1, days + 1):
        ws.cell(r, 2 + d, round(sum(ws.cell(5 + 3 * i + 1, 2 + d).value for i in range(3)), 1))
    ws.cell(r + 2, 1, "Remarks: CRM06 shut on 12th for maintenance")
    return wb_bytes(wb)


def stacked(year: int, month: int) -> bytes:
    """Two stacked sections with own headers; date strings / day numbers; MTD row; Yield %."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily"
    ws["A1"] = f"Daily Production Statement - {calendar.month_name[month]} {year}"
    ws["A3"] = "PICKLING LINE"
    ws.append([])
    ws["A4"], ws["B4"], ws["C4"], ws["D4"] = "Date", "Input", "Output", "Yield %"
    days = calendar.monthrange(year, month)[1]
    r = 5
    ti = to = 0
    for d in range(1, days + 1):
        ws.cell(r, 1, f"{d:02d}.{month:02d}.{year}")
        i, o = val(1, d, 1200), val(2, d, 1150)
        ws.cell(r, 2, i)
        ws.cell(r, 3, o)
        ws.cell(r, 4, round(o / i * 100, 2))
        ti += i
        to += o
        r += 1
    ws.cell(r, 1, "Total")
    ws.cell(r, 2, round(ti, 1))
    ws.cell(r, 3, round(to, 1))
    r += 3
    ws.cell(r, 1, "CRM04")
    r += 1
    ws.cell(r, 1, "Day")
    ws.cell(r, 2, "Input")
    ws.cell(r, 3, "Output")
    ws.cell(r, 4, "Delay (Hrs)")
    ws.cell(r, 5, "Cum Output")
    r += 1
    cum = 0
    for d in range(1, days + 1):
        o = val(3, d, 900)
        cum += o
        ws.cell(r, 1, d)
        ws.cell(r, 2, val(4, d, 950))
        ws.cell(r, 3, o)
        ws.cell(r, 4, dt.time(1, 30) if d == 5 else (d % 4))
        ws.cell(r, 5, round(cum, 1))
        r += 1
    ws.cell(r, 1, "MTD")
    ws.cell(r, 3, round(cum, 1))
    return wb_bytes(wb)


def long_format(year: int, month: int) -> bytes:
    """Date (merged over 3 shift rows) | Mill | Shift | Input | Output."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shiftwise"
    ws.append(["Shift wise production", None, None, None, None])
    ws.append(["Date", "Mill", "Shift", "Input (MT)", "Output (MT)"])
    r = 3
    for d in range(1, 11):
        ws.cell(r, 1, dt.datetime(year, month, d))
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=1)
        for k, sh in enumerate("ABC"):
            ws.cell(r + k, 2, "CRM04")
            ws.cell(r + k, 3, sh)
            ws.cell(r + k, 4, val(k, d, 300))
            ws.cell(r + k, 5, val(k + 1, d, 290))
        r += 3
    # unlabelled continuation row (no date, no label) -> must go to unmapped
    ws.cell(r, 4, 123.0)
    ws.cell(r, 5, 120.0)
    return wb_bytes(wb)


def monthly_summary(year: int, month: int) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Monthly Summary {calendar.month_abbr[month]}'{str(year)[2:]}"
    ws.append([])
    ws.append(["Mill", "Plan (MT)", "Actual (MT)", "Achievement %"])
    rows = [("CRM04", 30000, 28750), ("CRM06", 25000, 26010), ("BAF", 20000, 18000)]
    for m, p, a in rows:
        ws.append([m, p, a, round(a / p * 100, 1)])
    ws.append(["Total", sum(x[1] for x in rows), sum(x[2] for x in rows), None])
    return wb_bytes(wb)


def plant_template(year: int, month: int) -> bytes:
    """Mimics the CRM 'PRODUCTION' sheet: merged title, 3-level header, helper day column,
    mill columns, both-mill total, product-segment breakdown, a 2-column 'ROLLING O/P' sum column
    whose formula is missing on the last day, an unlabelled copy column, TOTAL row after blank
    rows, repeated header rows and a calculation block below. All values are synthetic."""
    from openpyxl.utils import get_column_letter as L
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PRODUCTION"
    ws["A1"] = f"{calendar.month_name[month].upper()}-PRODUCTION NARROW-{year}"
    ws.merge_cells("A1:N1")
    hdr2 = {1: "DATE", 2: "CRM04", 6: "CRM06", 10: "BOTH MILL TOTAL O/P PRODUCTION", 13: "HNT", 16: "TUBE",
            19: "TUBE ROLLING O/P"}
    for c, v in hdr2.items():
        ws.cell(2, c, v)
    ws.merge_cells("B2:E2"); ws.merge_cells("F2:I2"); ws.merge_cells("J2:K2"); ws.merge_cells("M2:N2")
    ws.merge_cells("P2:Q2")
    for c, v in {2: "ROLLING", 4: "R/R", 6: "ROLLING", 8: "R/R", 13: "CRM04", 14: "CRM06", 16: "CRM04", 17: "CRM06"}.items():
        ws.cell(3, c, v)
    ws.merge_cells("B3:C3"); ws.merge_cells("D3:E3"); ws.merge_cells("F3:G3"); ws.merge_cells("H3:I3")
    for c in (2, 4, 6, 8):
        ws.cell(4, c, "Input")
        ws.cell(4, c + 1, "Output")
    ws.cell(4, 10, "o/p of both mill")
    ws.cell(4, 11, "as per sap")
    for c in (13, 14, 16, 17):
        ws.cell(4, c, "O/P")
    days = calendar.monthrange(year, month)[1]
    first = 6
    for d in range(1, 32):
        r = first + d - 1
        ws.cell(r, 12, d)                                          # helper day column L (1..31 always)
        if d > days:
            ws.cell(r, 10, 0)
            continue
        ws.cell(r, 1, dt.datetime(year, month, d))
        a, b, c_, e = val(1, d, 120), val(2, d, 40), val(3, d, 150), val(4, d, 50)
        ws.cell(r, 2, a + 1); ws.cell(r, 3, a)
        ws.cell(r, 4, b + 1); ws.cell(r, 5, b)
        ws.cell(r, 6, c_ + 1); ws.cell(r, 7, c_)
        ws.cell(r, 8, e + 1); ws.cell(r, 9, e)
        both = round(a + b + c_ + e, 3)
        ws.cell(r, 10, both)
        ws.cell(r, 11, both)
        # product segments: breakdown of the same mill output
        ws.cell(r, 13, a); ws.cell(r, 14, c_); ws.cell(r, 16, b); ws.cell(r, 17, e)
        ws.cell(r, 19, 0 if d == days else round(b + e, 3))    # TUBE sum col, formula "missing" on last day
        ws.cell(r, 21, both)                                     # unlabelled copy of BOTH MILL
    tot = first + 31
    ws.cell(tot, 1, "TOTAL")
    for c in list(range(2, 12)) + [13, 14, 16, 17, 19, 21]:
        ws.cell(tot, c, round(sum((ws.cell(r, c).value or 0) for r in range(first, first + days)), 3))
    for c in (2, 4, 6, 8):
        ws.cell(tot + 2, c, "Input")
        ws.cell(tot + 2, c + 1, "Output")
    ws.cell(tot + 3, 2, "ROLLING"); ws.cell(tot + 3, 4, "R/R")
    ws.cell(tot + 4, 2, "CRM04"); ws.cell(tot + 4, 6, "CRM06"); ws.cell(tot + 4, 13, "HNT")
    ws.cell(tot + 8, 1, "PRODUCTION OF CRM04")
    ws.cell(tot + 9, 1, "ROLLING"); ws.cell(tot + 9, 2, ws.cell(tot, 3).value); ws.cell(tot + 9, 3, "MT")
    return wb_bytes(wb)


def main(out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    apr = vertical_merged(2023, 4)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("Production/2023/Production_April_2023.xlsx", apr)
        z.writestr("Production/2023/Production_May_2023.xlsx",
                   vertical_merged(2023, 5, extra_pickling=True, unmerged=True, formula_total=True))
        z.writestr("Production/2024/June-24 Production.xlsx", horizontal(2024, 6))
        z.writestr("Production/2024/Jul_2024_stacked.xlsx", stacked(2024, 7))
        z.writestr("Production/2024/Aug 2024 shiftwise.xlsx", long_format(2024, 8))
        z.writestr("Production/2024/Sep-2024 summary.xlsx", monthly_summary(2024, 9))
        z.writestr("Production/plant/FEB-PRODUCTION NARROW-2025.xlsx", plant_template(2025, 2))
        z.writestr("Production/plant/MARCH-PRODUCTION NARROW-2025.xlsx", plant_template(2025, 3))
        z.writestr("Production/2023/copy/Production_April_2023 (1).xlsx", apr)  # duplicate file
        z.writestr("Production/2024/broken_Oct_2024.xlsx", b"this is not a real xlsx file")
        z.writestr("Production/2022/old_format_Mar_2022.xls", b"\xd0\xcf\x11\xe0legacy")
        z.writestr("Production/2023/~$Production_April_2023.xlsx", b"lock")
        z.writestr("__MACOSX/Production/._x.xlsx", b"meta")
        inner = io.BytesIO()
        with zipfile.ZipFile(inner, "w") as zi:
            zi.writestr("Nov_2024.xlsx", vertical_merged(2024, 11))
        z.writestr("Production/2024/archive_nov.zip", inner.getvalue())
    print(f"wrote {out}")


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else "tests/data/sample_production.zip"))
